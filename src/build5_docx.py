# -*- coding: utf-8 -*-
import build5 as B
from build5 import BUTIE, E, C, sections, L, TITLE, answer_lines, NOTE_ANS, card_fields
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
FONT="Microsoft JhengHei"
def setfont(run_or_style, size=None, bold=None, color=None):
    f=run_or_style.font
    f.name=FONT
    el=run_or_style.element if hasattr(run_or_style,'element') else None
    rpr=run_or_style._element.get_or_add_rPr() if hasattr(run_or_style,'_element') and hasattr(run_or_style._element,'get_or_add_rPr') else run_or_style.element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:
        rf=rpr.makeelement(qn('w:rFonts'),{}); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),FONT)
    if size: f.size=Pt(size)
    if bold is not None: f.bold=bold
    if color: f.color.rgb=RGBColor.from_string(color)
def make(color,path):
    d=Document()
    sec=d.sections[0]
    sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
    for m in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(sec,m,Cm(2.0))
    st=d.styles['Normal']; setfont(st,11); st.paragraph_format.space_after=Pt(3); st.paragraph_format.line_spacing=1.2
    TEAL="1F5C5A" if color else "000000"
    for n,sz in (('Heading 1',16),('Heading 2',14),('Heading 3',12)):
        hs=d.styles[n]; setfont(hs,sz,True,TEAL); hs.paragraph_format.space_before=Pt(12); hs.paragraph_format.space_after=Pt(4); hs.paragraph_format.keep_with_next=True
    def para(text="",bold=False,size=None,col=None,indent=0,align=None,first=None,keep=False):
        p=d.add_paragraph()
        if indent: p.paragraph_format.left_indent=Cm(indent)
        if first is not None: p.paragraph_format.first_line_indent=Cm(first)
        if align: p.alignment=align
        if keep: p.paragraph_format.keep_together=True; p.paragraph_format.keep_with_next=True
        if text:
            r=p.add_run(text); setfont(r,size,bold,col)
        return p
    def labeled(label,text,lc=None,indent=0.5):
        p=d.add_paragraph(); p.paragraph_format.left_indent=Cm(indent); p.paragraph_format.keep_together=True
        r=p.add_run(label+"："); setfont(r,None,True,lc if color else None)
        r=p.add_run(text); setfont(r)
        return p
    LC={"相似":"1F4E9E","相反":"C55A11","辨析":"7030A0","典故":"7030A0","歇後語":"7030A0","猜謎":"7030A0","俗諺":"7030A0","補充":"7030A0","備註":"7F7F7F"}
    p=para(TITLE,True,20,TEAL,align=WD_ALIGN_PARAGRAPH.CENTER)
    para("班級：＿＿＿＿　座號：＿＿＿＿　姓名：＿＿＿＿＿＿　得分：＿＿＿＿")
    d.add_heading("壹、成語學習站",1)
    for i,e in enumerate(E,1):
        para("%d %s"%(i,e['name']),True,13,TEAL if color else None,keep=True)
        for k,v,c in card_fields(e): labeled(k,v,LC.get(k))
    d.add_heading("★ 老師補充成語",2)
    for x in C.EXTRA:
        para("★ "+x['name'],True,13,"C55A11" if color else None,keep=True)
        labeled("解釋",x['jie'])
        if x['xs']: labeled("相似",x['xs'],LC["相似"])
        labeled("備註",x['note'],LC["備註"])
    d.add_heading("成語大補帖："+BUTIE[0],2)
    for k,v in BUTIE[1]: labeled(k,v,None)
    d.add_page_break()
    d.add_heading("貳、實力挑戰",1)
    for title,note,qs in sections:
        d.add_heading(title,2)
        if note: para(note,size=10,col="7F7F7F" if color else None)
        if title.startswith("五"):
            for t in C.PASSAGE: para(t,first=0.8)
        for i,q in enumerate(qs,1):
            if q['k']=='mc':
                para("（　）%d. %s"%(i,q['stem']),keep=True)
                if all(len(o)<=8 for o in q['opts']):
                    para("　".join("(%s) %s"%(L[j],o) for j,o in enumerate(q['opts'])),indent=0.8)
                else:
                    for j,o in enumerate(q['opts']):
                        p=para("(%s) %s"%(L[j],o),indent=0.8)
                        if j<3: p.paragraph_format.keep_with_next=True
            elif q['k']=='fill':
                para("%d. %s"%(i,q['tpl'].replace('{}','（　）')))
            else:
                para("%d. %s → （　　　　　　　）"%(i,q['stem']))
    d.add_page_break()
    d.add_heading("參考答案",1)
    for t,s in answer_lines():
        p=para(indent=0); r=p.add_run(t+"："); setfont(r,None,True,TEAL if color else None); r=p.add_run(s); setfont(r)
    para(NOTE_ANS,size=10,col="7F7F7F" if color else None)
    for z in d.settings.element.iter(qn('w:zoom')):
        z.set(qn('w:percent'),'100')
    d.core_properties.title="成語習作 第5回"; d.core_properties.author=""
    d.save(path)
make(False,B.OUT+"/xizuo-5-handout.docx")
make(True,B.OUT+"/xizuo-5-handout-color.docx")
print("ok")
